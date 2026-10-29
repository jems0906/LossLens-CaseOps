from pathlib import Path
import csv
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "fixtures"
OUTPUT = ROOT / "docs" / "excel-proof" / "losslens_synthetic_practice.xlsx"
EVENT_FILES = [
    "synthetic_inventory_events.csv",
    "synthetic_transfer_events.csv",
    "synthetic_return_events.csv",
    "synthetic_adjustment_events.csv",
]


def load_events() -> list[dict[str, str]]:
    rows = []
    for filename in EVENT_FILES:
        with (FIXTURES / filename).open(encoding="utf-8-sig", newline="") as source:
            rows.extend(csv.DictReader(source))
    return rows


def style_sheet(sheet, table_name: str | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor="274936")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
        sheet.column_dimensions[letter].width = width
    if table_name and sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Start Here"
    guide.append(["LossLens CaseOps | Synthetic Excel Practice"])
    guide.append(["Dataset", "Fictional shipment, receipt, transfer, return, and inventory adjustment records only."])
    guide.append(["Exercise status", "Workbook scaffold generated. Build and screenshot actual PivotTables in desktop Excel."])
    guide.append(["Duplicate check", "Events table: =COUNTIFS(Events[event_id],[@event_id])>1"])
    guide.append(["Expected units", "Reconciliation table uses SUMIFS by reconciliation_id and event_type."])
    guide.append(["Recorded units", "Receipt, transfer_out, and return rows are recorded at this defined handoff step."])
    guide.append(["Missing values", "Keep missing measures blank/unknown. Never coerce an absent denominator to zero."])
    guide.append(["Timezone", "Source timestamps are UTC ISO 8601. Preserve raw values; normalize before grouping by week."])
    guide.append(["PivotTables", "Insert PivotTables from Events/Cases. Group by fictional site, process, and normalized week."])
    guide.append(["Caution", "A records mismatch is not proof of theft or misconduct."])
    style_sheet(guide)
    guide.column_dimensions["A"].width = 23
    guide.column_dimensions["B"].width = 105
    guide.freeze_panes = "A2"

    events_sheet = workbook.create_sheet("Events")
    event_headers = list(load_events()[0].keys()) + ["duplicate_id", "normalized_week"]
    events_sheet.append(event_headers)
    for row_number, event in enumerate(load_events(), start=2):
        events_sheet.append([event[key] for key in event_headers[:8]] + [
            f'=COUNTIFS(Events[event_id],A{row_number})>1',
            f'=TEXT(DATEVALUE(LEFT(E{row_number},10))-WEEKDAY(DATEVALUE(LEFT(E{row_number},10)),2)+1,"yyyy-mm-dd")',
        ])
    style_sheet(events_sheet, "Events")
    events_sheet.column_dimensions["E"].width = 30

    cases = json.loads((FIXTURES / "synthetic_cases.json").read_text(encoding="utf-8"))
    case_sheet = workbook.create_sheet("Cases")
    case_fields = ["case_id", "simulated", "case_type", "report_source", "priority", "owner", "status", "site_id", "process_step", "reconciliation_id", "expected_units", "recorded_units", "raw_discrepancy_units", "unexplained_units", "opened_at", "closed_at"]
    case_sheet.append(case_fields + ["open_age_days", "closed_duration_days"])
    for row_number, case in enumerate(cases, start=2):
        case_sheet.append([case.get(field) for field in case_fields] + [
            f'=IF(G{row_number}="closed","",TODAY()-DATEVALUE(LEFT(O{row_number},10)))',
            f'=IF(G{row_number}<>"closed","",DATEVALUE(LEFT(P{row_number},10))-DATEVALUE(LEFT(O{row_number},10)))',
        ])
    style_sheet(case_sheet, "Cases")

    reconciliation = workbook.create_sheet("Reconciliation")
    reconciliation.append(["reconciliation_id", "site_id", "process_step", "week_start", "expected_units", "recorded_units", "correction_units", "documented_adjustment_units", "correction_source_ref", "raw_mismatch", "expected_minus_recorded", "unexplained_units", "discrepancy_rate"])
    unique = {}
    for event in load_events():
        unique[event["reconciliation_id"]] = (event["site_id"], event["process_step"], event["occurred_at"][:10])
    for row_number, (reconciliation_id, (site, process, event_date)) in enumerate(sorted(unique.items()), start=2):
        reconciliation.append([
            reconciliation_id,
            site,
            process,
            f'=DATEVALUE("{event_date}")-WEEKDAY(DATEVALUE("{event_date}"),2)+1',
            f'=SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"shipment")+SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"transfer_in")',
            f'=SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"receipt")+SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"transfer_out")+SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"return")',
            f'=SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"receipt_correction")',
            f'=SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"damage_adjustment")+SUMIFS(Events[quantity],Events[reconciliation_id],A{row_number},Events[event_type],"inventory_adjustment")',
            f'=XLOOKUP(1,(Events[reconciliation_id]=A{row_number})*(Events[event_type]="receipt_correction"),Events[source_ref],"No correction record")',
            f'=ABS(E{row_number}-F{row_number})',
            f'=E{row_number}-F{row_number}',
            f'=IF(K{row_number}>0,MAX(0,J{row_number}-MIN(J{row_number},G{row_number}+H{row_number})),J{row_number})',
            f'=IF(E{row_number}=0,"",L{row_number}/E{row_number})',
        ])
    style_sheet(reconciliation, "Reconciliation")
    reconciliation.column_dimensions["D"].number_format = "yyyy-mm-dd"
    reconciliation.column_dimensions["M"].number_format = "0.00%"

    summary = workbook.create_sheet("One Page Summary")
    summary.append(["Measure", "Value / evidence", "Limit / next step"])
    summary.append(["CASE-SYN-001 first pass", "100 expected; 92 recorded; 8 raw mismatch", "Not a theft conclusion"])
    summary.append(["CASE-SYN-001 review", "5 delayed correction + 3 documented damage = 8 explained", "Resolved reconciliation issue; 0 unexplained"])
    summary.append(["CASE-SYN-002", "12 expected; 10 recorded; 2 unresolved", "Keep open; request synthetic correction and handoff records"])
    summary.append(["Process control", "Timestamped two-party handoff count and linked correction reason", "Synthetic receiving process owner; due 2026-11-15"])
    summary.append(["Rate definition", "Unexplained units / expected units", "Only compare same site, process, and defined period"])
    summary.append(["Time to closure", "=MEDIAN(Cases[closed_duration_days])", "Closed-case sample size: =COUNT(Cases[closed_duration_days])"])
    summary.append(["Action completion", "Use COUNTIFS for verified completed and due actions", "Leave blank if no actions were due; require verification evidence"])
    summary.append(["Limitations", "Synthetic records; one closed-case example", "Hashes do not prove authenticity; application audit is not immutable"])
    style_sheet(summary)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 57
    summary.column_dimensions["C"].width = 63
    for row in summary.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()